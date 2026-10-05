"""tests/test_fallback_extractor.py

Tests for the Flexible / Dynamic Extraction Engine:

* fuzzy column mapping (``formatter.build_column_mapping`` / currency detection);
* header-row auto-detection (``extractor_pdf.detect_header_row``);
* the Smart Fallback Extractor (``extractor_pdf.smart_extract_tables`` /
  ``smart_fallback_extract``) that activates for unknown suppliers;
* seamless export so the live EGP-converter formulas are still generated for
  unknown suppliers (``exporter``);
* the end-to-end pipeline (``main.run``) with an unknown supplier forced.

Runnable directly::

    python tests/test_fallback_extractor.py

or with pytest (``pytest tests/test_fallback_extractor.py``).
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

# Make the project root importable when run directly as a script.
ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import exporter  # noqa: E402
import extractor_pdf  # noqa: E402
import formatter  # noqa: E402
import main  # noqa: E402


STANDARD_COLUMNS = list(formatter.load_rules()["standard_columns"])


# ---------------------------------------------------------------------------
# 1. Fuzzy column mapping
# ---------------------------------------------------------------------------
def test_date_aliases_map_to_date():
    """``Dt`` / ``Date of Service`` / ``التاريخ`` all resolve to ``Date``."""
    for header in ["Dt", "DT.", "date", "Date of Service", "service date", "التاريخ"]:
        mapping = formatter.build_column_mapping([header])
        assert mapping.get(header) == "Date", f"{header!r} -> {mapping}"


def test_other_aliases_map_to_standard_columns():
    cases = {
        "Hotel Name": "Hotel",
        "الفندق": "Hotel",
        "Voucher No": "V.NO",
        "pax": "Adult",
        "Ad.": "Adult",
        "ch": "Child",
        "inf": "Inf",
        "Cost": "Total",
        "Net": "Net",
    }
    for header, expected in cases.items():
        mapping = formatter.build_column_mapping([header])
        assert mapping.get(header) == expected, (
            f"{header!r} -> {mapping} (expected {expected})"
        )


def test_currency_symbols_map_to_currency_columns():
    """``$``, ``£``, ``€``, ``EGP`` and ``L.E`` headers map to currency columns."""
    cases = {
        "$0": "USD",
        "£0": "GBP",
        "€ 0": "Euros",
        "EGP -": "EGP",
        "cash L.E": "EGP",
        "REST $": "USD",
        "egyptian pounds": "EGP",
    }
    for header, expected in cases.items():
        assert formatter.detect_header_currency(header) == expected, header
        mapping = formatter.build_column_mapping([header])
        assert mapping.get(header) == expected, f"{header!r} -> {mapping}"


def test_cost_header_with_symbol_is_total():
    """A header naming a cost (``Cost $``) becomes ``Total``, not ``USD``."""
    assert formatter.detect_header_currency("Cost $") is None
    assert formatter.build_column_mapping(["Cost $"])["Cost $"] == "Total"


def test_unknown_header_is_not_mapped():
    mapping = formatter.build_column_mapping(["Zorblax Field"])
    assert "Zorblax Field" not in mapping


def test_no_standard_column_claimed_twice():
    """Even with several similar headers, each standard column is claimed once."""
    headers = ["Dt", "Date of Service", "التاريخ"]
    mapping = formatter.build_column_mapping(headers)
    assert list(mapping.values()).count("Date") == 1


# ---------------------------------------------------------------------------
# 2. Header-row auto-detection (text density)
# ---------------------------------------------------------------------------
def test_detect_header_row_skips_title_and_numeric_rows():
    table = [
        ["ACME Travel  -  Invoice 09/2026", "", "", ""],
        ["Date", "Hotel", "Adult", "Cost"],
        ["1-Sep", "Hurghada", "2", "100"],
        ["Total", "", "", "100"],
    ]
    assert extractor_pdf.detect_header_row(table) == 1


def test_detect_header_row_on_empty_table():
    assert extractor_pdf.detect_header_row([]) == 0


# ---------------------------------------------------------------------------
# 3. Smart Fallback Extractor (mock unknown-supplier tables)
# ---------------------------------------------------------------------------
def _mock_tables():
    """A mock unknown-supplier invoice: fuzzy headers + a summary block."""
    transactions = [
        ["Dt", "Hotel", "Ad.", "ch", "inf", "$0", "EGP -", "Cost"],
        ["1-Sep", "Hurghada", "2", "1", "0", "100", "0", "100"],
        ["2-Sep", "Cairo", "1", "0", "0", "0", "500", "500"],
        ["Total", "", "", "", "", "", "", "600"],
    ]
    summary = [["Total Invoice", "600"], ["Net Invoice", "600"]]
    return [transactions, summary]


def test_smart_extract_tables_maps_and_drops_summary():
    formatted, summary = extractor_pdf.smart_extract_tables(_mock_tables())

    # Every standard column is present, in schema order.
    assert list(formatted.columns) == STANDARD_COLUMNS
    # The summary row is separated out; two data rows remain.
    assert len(formatted) == 2

    assert list(formatted["Hotel"]) == ["Hurghada", "Cairo"]
    assert list(formatted["Adult"]) == [2, 1]
    assert list(formatted["Child"]) == [1, 0]
    assert list(formatted["Inf"]) == [0, 0]
    assert list(formatted["USD"]) == [100, 0]
    assert list(formatted["EGP"]) == [0, 500]
    assert list(formatted["Total"]) == [100, 500]
    # ``Date`` is normalized to ``YYYY-MM-DD``.
    assert all(d.startswith("20") and d.count("-") == 2 for d in formatted["Date"])
    # Summary metrics parsed from the summary block.
    assert summary.get("Total Invoice") == 600


def test_smart_extract_tables_maps_arabic_headers():
    tables = [
        [
            ["التاريخ", "الفندق", "بالغ", "طفل"],
            ["1-Sep", "الغردقة", "2", "1"],
            ["2-Sep", "القاهرة", "1", "0"],
        ]
    ]
    formatted, _ = extractor_pdf.smart_extract_tables(tables)
    assert list(formatted.columns) == STANDARD_COLUMNS
    assert len(formatted) == 2
    assert list(formatted["Adult"]) == [2, 1]
    assert list(formatted["Child"]) == [1, 0]


def test_smart_extract_tables_empty_when_no_data():
    formatted, summary = extractor_pdf.smart_extract_tables([])
    assert list(formatted.columns) == STANDARD_COLUMNS
    assert formatted.empty
    assert summary == {}


# ---------------------------------------------------------------------------
# 4. Export: live EGP-converter formulas are still generated
# ---------------------------------------------------------------------------
def test_fallback_export_has_live_egp_formulas():
    formatted, summary = extractor_pdf.smart_extract_tables(_mock_tables())

    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "unknown_supplier.xlsx"
        exporter.export_to_excel(formatted, out, summary=summary)

        wb = load_workbook(out)
        ws = wb["Invoices"]
        header_row = exporter.HEADER_ROW
        headers = [cell.value for cell in ws[header_row]]
        assert "EGP Converter" in headers

        conv_index = headers.index("EGP Converter") + 1
        conv_cell = ws.cell(row=header_row + 1, column=conv_index)
        assert isinstance(conv_cell.value, str)
        assert conv_cell.value.startswith("=")
        # The formula must reference the editable USD/EUR/GBP rate cells.
        assert "$B$2" in conv_cell.value  # USD rate
        assert "$D$2" in conv_cell.value  # EUR rate
        assert "$F$2" in conv_cell.value  # GBP rate


# ---------------------------------------------------------------------------
# 4b. Subtotal ("Total") row
# ---------------------------------------------------------------------------
def test_export_appends_total_row_with_sum_formulas():
    formatted, _ = extractor_pdf.smart_extract_tables(_mock_tables())
    n = len(formatted)
    assert n == 2

    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "with_total.xlsx"
        exporter.export_to_excel(formatted, out)
        ws = load_workbook(out)["Invoices"]

        header_row = exporter.HEADER_ROW
        headers = {str(c.value): c.column_letter for c in ws[header_row]}
        data_start = header_row + 1
        data_end = header_row + n
        total_row = data_end + 1

        # The label sits in a text column (Hotel is preferred).
        assert ws[f"{headers['Hotel']}{total_row}"].value == "Total"

        # Every financial column gets a SUM over the exact data range.
        for name in ["EGP Converter", "EGP", "GBP", "Euros", "USD", "Total", "Net"]:
            letter = headers[name]
            assert ws[f"{letter}{total_row}"].value == (
                f"=SUM({letter}{data_start}:{letter}{data_end})"
            ), name

        # Per-row EGP-converter formulas remain intact (not overwritten).
        conv = headers["EGP Converter"]
        for row in range(data_start, data_end + 1):
            value = ws[f"{conv}{row}"].value
            assert isinstance(value, str) and value.startswith("=")
            assert "$B$2" in value and "$D$2" in value and "$F$2" in value

        # Subtotal styling: bold + a double top border.
        total_cell = ws[f"{headers['USD']}{total_row}"]
        assert total_cell.font.bold is True
        assert total_cell.border.top.style == "double"

        # The AutoFilter still covers only the data rows (no subtotal row).
        assert ws.auto_filter.ref.endswith(str(data_end))
        assert str(total_row) not in ws.auto_filter.ref


def test_export_total_row_label_is_configurable():
    formatted, _ = extractor_pdf.smart_extract_tables(_mock_tables())
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "arabic_total.xlsx"
        exporter.export_to_excel(formatted, out, total_label="الاجمالي")
        ws = load_workbook(out)["Invoices"]

        header_row = exporter.HEADER_ROW
        headers = {str(c.value): c.column_letter for c in ws[header_row]}
        total_row = header_row + len(formatted) + 1
        assert ws[f"{headers['Hotel']}{total_row}"].value == "الاجمالي"


def test_export_empty_df_has_no_total_row():
    formatted, _ = extractor_pdf.smart_extract_tables([])
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "empty.xlsx"
        exporter.export_to_excel(formatted, out)
        ws = load_workbook(out)["Invoices"]
        # Only the header row exists below the rate block; no subtotal appended.
        assert ws.max_row == exporter.HEADER_ROW


# ---------------------------------------------------------------------------
# 5. End-to-end pipeline with an unknown supplier
# ---------------------------------------------------------------------------
def test_detect_supplier_returns_none_for_unknown_text():
    assert formatter.detect_supplier("Totally Unknown Supplier Co.") is None


def test_main_run_falls_back_for_unknown_supplier():
    pdfs = sorted(ROOT.glob("*.pdf")) + sorted(ROOT.glob("*.PDF"))
    if not pdfs:
        print("      (skipped: no sample PDF in the workspace)")
        return

    pdf_path = pdfs[0]
    with tempfile.TemporaryDirectory() as tmp:
        out_path = Path(tmp) / "fallback_out.xlsx"
        out, formatted, summary, supplier_key = main.run(
            pdf_path, supplier_key="__unknown_supplier__", output_path=out_path
        )

        assert supplier_key == extractor_pdf.DYNAMIC_SUPPLIER_KEY
        assert list(formatted.columns) == STANDARD_COLUMNS
        assert out.exists()

        wb = load_workbook(out)
        headers = [cell.value for cell in wb["Invoices"][exporter.HEADER_ROW]]
        assert "EGP Converter" in headers


# ---------------------------------------------------------------------------
# 6. Hybrid mode: supplier profiles first, Smart Fallback when unknown
# ---------------------------------------------------------------------------
def test_main_run_uses_profile_for_known_supplier():
    """A known supplier uses its curated profile and is reported as such."""
    pdfs = sorted(ROOT.glob("*.pdf")) + sorted(ROOT.glob("*.PDF"))
    if not pdfs:
        print("      (skipped: no sample PDF in the workspace)")
        return

    known = sorted(formatter.load_rules().get("suppliers", {}))
    if not known:
        print("      (skipped: no supplier profiles configured)")
        return

    with tempfile.TemporaryDirectory() as tmp:
        out_path = Path(tmp) / "profile_forced.xlsx"
        forced = known[0]
        out, formatted, summary, supplier_key = main.run(
            pdfs[0], supplier_key=forced, output_path=out_path, engine="text"
        )

        assert supplier_key == forced, (
            f"expected the forced profile {forced!r}, got {supplier_key!r}"
        )
        assert main.supplier_label(supplier_key) == f"{forced} (profile)"
        assert list(formatted.columns) == STANDARD_COLUMNS
        assert out.exists()


def test_main_run_hybrid_routes_by_detection():
    """Known -> profile mapper, unknown -> Smart Fallback, never both."""
    pdfs = sorted(ROOT.glob("*.pdf")) + sorted(ROOT.glob("*.PDF"))
    if not pdfs:
        print("      (skipped: no sample PDF in the workspace)")
        return

    pdf_path = pdfs[0]
    detected = formatter.detect_supplier(
        extractor_pdf.get_pdf_text(pdf_path), formatter.load_rules()
    )
    expected_key = detected if detected else extractor_pdf.DYNAMIC_SUPPLIER_KEY

    calls = {"fallback": 0, "profile": 0}
    orig_fallback = extractor_pdf.smart_fallback_extract
    orig_tables = extractor_pdf.extract_tables_from_pdf

    def _spy_fallback(*args, **kwargs):
        calls["fallback"] += 1
        return orig_fallback(*args, **kwargs)

    def _spy_profile(*args, **kwargs):
        calls["profile"] += 1
        return orig_tables(*args, **kwargs)

    extractor_pdf.smart_fallback_extract = _spy_fallback
    extractor_pdf.extract_tables_from_pdf = _spy_profile
    try:
        with tempfile.TemporaryDirectory() as tmp:
            out, formatted, summary, supplier_key = main.run(
                pdf_path, output_path=Path(tmp) / "hybrid.xlsx", engine="text"
            )
    finally:
        extractor_pdf.smart_fallback_extract = orig_fallback
        extractor_pdf.extract_tables_from_pdf = orig_tables

    assert supplier_key == expected_key, (
        f"expected {expected_key!r} for detected={detected!r}, got {supplier_key!r}"
    )
    if detected:
        assert calls["profile"] == 1, "known supplier must use its profile"
        assert calls["fallback"] == 0, "known supplier must not use the fallback"
    else:
        assert calls["fallback"] == 1, "unknown supplier must use Smart Fallback"
        assert calls["profile"] == 0, "unknown supplier must not use a profile"


def test_forced_dynamic_key_uses_smart_fallback():
    """Passing the dynamic key explicitly forces Smart Fallback."""
    pdfs = sorted(ROOT.glob("*.pdf")) + sorted(ROOT.glob("*.PDF"))
    if not pdfs:
        print("      (skipped: no sample PDF in the workspace)")
        return

    with tempfile.TemporaryDirectory() as tmp:
        out, formatted, summary, supplier_key = main.run(
            pdfs[0],
            supplier_key=extractor_pdf.DYNAMIC_SUPPLIER_KEY,
            output_path=Path(tmp) / "forced_fallback.xlsx",
            engine="text",
        )

        assert supplier_key == extractor_pdf.DYNAMIC_SUPPLIER_KEY
        assert main.supplier_label(supplier_key) == "Unknown (Smart Fallback)"
        assert list(formatted.columns) == STANDARD_COLUMNS


def test_supplier_label_reports_strategy():
    """supplier_label() distinguishes profile vs Smart Fallback."""
    assert main.supplier_label("egypt_gate") == "egypt_gate (profile)"
    assert main.supplier_label("smile_safari") == "smile_safari (profile)"
    assert (
        main.supplier_label(extractor_pdf.DYNAMIC_SUPPLIER_KEY)
        == "Unknown (Smart Fallback)"
    )
    assert main.supplier_label(None) == "Unknown (Smart Fallback)"


def test_count_detected_columns_counts_populated_fields():
    """The detected-column count reflects populated standard columns only."""
    formatted = pd.DataFrame(
        {
            "Hotel": ["Hurghada", "Cairo"],
            "Date": ["2026-01-02", "2026-01-03"],
            "Adult": [2, 1],
            "Total": [np.nan, np.nan],  # unmapped -> not detected
            "Net": ["", "  "],  # blank strings -> not detected
        }
    )
    schema = ["Hotel", "Date", "Adult", "Total", "Net", "Missing"]
    assert main.count_detected_columns(formatted, schema) == 3


def test_count_detected_columns_empty_frame_is_zero():
    """An empty/None frame reports 0 rather than raising."""
    assert main.count_detected_columns(pd.DataFrame(columns=STANDARD_COLUMNS)) == 0
    assert main.count_detected_columns(None) == 0


def test_wide_sparse_table_keeps_its_rows():
    """Wide landscape tables must not be emptied by the sparsity filter.

    Regression: ``min_filled`` used to grow with the column count
    (``n_cols // 2``), so a 26-column invoice whose rows fill ~8 cells was held
    to 13 filled cells and EVERY row was discarded as a summary row. Because
    Smart Fallback now handles every invoice, that silently produced empty
    workbooks for such layouts.
    """
    header = ["DATE", "D", "S", "PRIVET", "2", "4 SEAT", "$0", "EGP -", "TIME",
              "HOTEL", "Room", "1", "CAMEL"]
    # Pad out to a wide table (26 columns) with mostly-empty cells.
    header = header + [f"spacer{i}" for i in range(26 - len(header))]
    row = ["01/09/2026", "", "", "EGYPT GATE", "2", "4", "100", "4500",
           "10:00", "CAMEL", "101", "2", "ADULTS"] + [None] * 13
    rows = [header] + [list(row) for _ in range(3)]

    raw, summary = extractor_pdf.tables_to_transactions([rows])

    assert not raw.empty, "wide sparse table lost every row"
    assert len(raw) == 3, f"expected 3 data rows, got {len(raw)}"


def test_narrow_table_still_drops_sparse_rows():
    """The narrow-table path keeps dropping genuinely sparse placeholder rows."""
    rows = [
        ["Date", "Hotel", "Adult", "Total"],
        ["2026-01-01", "Hurghada", "2", "3000"],
        [None, None, None, None],  # sparse placeholder -> dropped
    ]
    raw, _ = extractor_pdf.tables_to_transactions([rows])
    assert len(raw) == 1, f"expected only the populated row, got {len(raw)}"


# ---------------------------------------------------------------------------
# Standalone runner (no pytest required)
# ---------------------------------------------------------------------------
def _run_all():
    tests = [
        value
        for name, value in sorted(globals().items())
        if name.startswith("test_") and callable(value)
    ]
    passed = 0
    for test in tests:
        try:
            test()
            print(f"PASS  {test.__name__}")
            passed += 1
        except AssertionError as exc:
            print(f"FAIL  {test.__name__}: {exc}")
        except Exception as exc:  # noqa: BLE001 - surface any unexpected error
            print(f"ERROR {test.__name__}: {type(exc).__name__}: {exc}")
    print(f"\n{passed}/{len(tests)} tests passed")
    return 0 if passed == len(tests) else 1


if __name__ == "__main__":
    raise SystemExit(_run_all())
