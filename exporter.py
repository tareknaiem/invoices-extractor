"""exporter.py

Write a standardized DataFrame to a formatted Excel workbook using openpyxl.

Adds an editable exchange-rate block above the data, an AutoFilter, a frozen
header row, auto-fitted column widths, and sensible number formats. The
"EGP Converter" column is populated with live formulas that convert USD/EUR/GBP
amounts to EGP using the rate cells. An optional ``summary`` dict is written to
a separate "Summary" sheet.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# Standard columns that hold monetary values -> currency number format.
CURRENCY_COLUMNS = {"EGP Converter", "EGP", "GBP", "Euros", "USD", "Total", "Net"}
# Standard columns that hold whole-number counts.
COUNT_COLUMNS = {"Adult", "Child", "Inf"}

# Default exchange rates (EGP per 1 unit of foreign currency).
EXCHANGE_RATES = {"USD": 50.0, "EUR": 55.0, "GBP": 65.0}
# Cell references where each exchange-rate input lives (see ``_write_rate_block``).
RATE_CELLS = {"USD": "$B$2", "EUR": "$D$2", "GBP": "$F$2"}
# The data-table header row (rows 1-3 hold the exchange-rate block).
HEADER_ROW = 4

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
DATA_ALIGNMENT = Alignment(vertical="center")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
RATE_CELL_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
HINT_FONT = Font(italic=True, color="808080")


def _style_sheet(ws, currency_columns=(), count_columns=(), header_row=HEADER_ROW):
    """Style a data sheet: header, AutoFilter, freeze, borders, number formats."""
    max_row = ws.max_row
    max_col = ws.max_column
    if max_col == 0:
        return

    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(max_col)}{max_row}"
    ws.freeze_panes = f"A{header_row + 1}"

    header_map = {}  # column name -> column letter
    for cell in ws[header_row]:
        if cell.value is not None:
            header_map[str(cell.value)] = cell.column_letter
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    for row in ws.iter_rows(min_row=header_row + 1, max_row=max_row, max_col=max_col):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = DATA_ALIGNMENT

    for col_name in currency_columns:
        letter = header_map.get(col_name)
        if not letter:
            continue
        for row in range(header_row + 1, max_row + 1):
            ws[f"{letter}{row}"].number_format = "#,##0.00"

    for col_name in count_columns:
        letter = header_map.get(col_name)
        if not letter:
            continue
        for row in range(header_row + 1, max_row + 1):
            ws[f"{letter}{row}"].number_format = "0"


def _autofit_columns(ws, min_width=10, max_width=40, start_row=1):
    """Approximate column widths from the longest visible cell value.

    ``start_row`` skips any header/summary block above the data (e.g. the
    exchange-rate block on the Invoices sheet).
    """
    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in col:
            if cell.row < start_row:
                continue
            value = cell.value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        width = max(min_width, min(max_width, max_len + 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _style_summary(ws):
    """Style the optional summary sheet."""
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autofit_columns(ws)
    for row in range(2, ws.max_row + 1):
        ws[f"B{row}"].number_format = "#,##0.00"


def _write_rate_block(ws):
    """Write the editable exchange-rate block above the data (rows 1-3)."""
    ws["A1"] = "Exchange Rates (EGP)"
    ws["A1"].font = Font(bold=True, size=12)

    # Row 2: label + editable rate cell for USD, EUR, GBP.
    column = 1
    for currency in ("USD", "EUR", "GBP"):
        label = ws.cell(row=2, column=column)
        label.value = currency
        label.font = Font(bold=True)

        rate = ws.cell(row=2, column=column + 1)
        rate.value = EXCHANGE_RATES[currency]
        rate.number_format = "#,##0.00"
        rate.fill = RATE_CELL_FILL
        rate.border = THIN_BORDER
        column += 2

    # Row 3: hint.
    ws["A3"] = "Edit the highlighted cells to update conversion rates."
    ws["A3"].font = HINT_FONT


def _write_converter_formulas(ws, n_rows):
    """Populate the "EGP Converter" column with live EGP-conversion formulas.

    Each row computes: EGP + USD*USD_rate + Euros*EUR_rate + GBP*GBP_rate.
    ``N()`` coerces empty/text cells to 0 so blank values don't break the math.
    """
    header_map = {cell.value: cell.column_letter for cell in ws[HEADER_ROW]}
    conv_col = header_map.get("EGP Converter")
    egp_col = header_map.get("EGP")
    usd_col = header_map.get("USD")
    eur_col = header_map.get("Euros")
    gbp_col = header_map.get("GBP")

    if not all([conv_col, egp_col, usd_col, eur_col, gbp_col]):
        return

    for i in range(n_rows):
        row = HEADER_ROW + 1 + i
        formula = (
            f"=N({egp_col}{row})"
            f"+N({usd_col}{row})*{RATE_CELLS['USD']}"
            f"+N({eur_col}{row})*{RATE_CELLS['EUR']}"
            f"+N({gbp_col}{row})*{RATE_CELLS['GBP']}"
        )
        cell = ws[f"{conv_col}{row}"]
        cell.value = formula
        cell.number_format = "#,##0.00"


def export_to_excel(df, output_path, summary=None, sheet_name="Invoices"):
    """Write ``df`` (and an optional ``summary`` dict) to a formatted workbook.

    The "Invoices" sheet gets an editable exchange-rate block at the top and a
    "EGP Converter" column filled with live EGP-conversion formulas.

    Args:
        df (pd.DataFrame): The standardized transactions to export.
        output_path (str or Path): Destination ``.xlsx`` path.
        summary (dict, optional): Key/value summary metrics written to a
            separate "Summary" sheet.
        sheet_name (str): Name of the transactions sheet.

    Returns:
        Path: The path to the written workbook.
    """
    output_path = Path(output_path)
    if output_path.suffix.lower() != ".xlsx":
        output_path = output_path.with_suffix(".xlsx")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # Write the data table starting below the exchange-rate block (header row 4).
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(
            writer,
            sheet_name=sheet_name,
            index=False,
            na_rep="",
            startrow=HEADER_ROW - 1,
        )
        if summary:
            summary_df = pd.DataFrame(list(summary.items()), columns=["Metric", "Value"])
            summary_df.to_excel(writer, sheet_name="Summary", index=False)

    wb = load_workbook(output_path)
    ws = wb[sheet_name]

    _write_rate_block(ws)
    _write_converter_formulas(ws, len(df))
    _style_sheet(ws, CURRENCY_COLUMNS, COUNT_COLUMNS)
    _autofit_columns(ws, start_row=HEADER_ROW)

    if summary and "Summary" in wb.sheetnames:
        _style_summary(wb["Summary"])

    wb.save(output_path)
    return output_path
